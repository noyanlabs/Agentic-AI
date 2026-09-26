# Reads an Excel workbook, loads every sheet into SQLite and gives the LLM the schema + properties so it can write SQL queries.
from pathlib import Path

import pandas as pd

from ProjectBackend.Interpreter import empty_result, error_result
from ProjectBackend.Interpreter.sql_utils import build_database, describe_database, schema_text

KIND = "excel"


def _sheet_properties(path: Path) -> dict:
    # workbook-level properties that plain DataFrames lose: merged cells, formulas, hidden sheets
    props = {}
    if path.suffix.lower() == ".xls":
        return props
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False, read_only=False)
        for ws in wb.worksheets:
            formulas = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
            props[ws.title] = {"visible": ws.sheet_state == "visible", "merged_ranges": len(ws.merged_cells.ranges), "formulas": formulas,
                               "dimensions": ws.dimensions, "charts": len(getattr(ws, "_charts", []))}
        wb.close()
    except Exception:
        pass
    return props


def initialize(path, **options) -> dict:
    p = Path(path)
    try:
        frames = pd.read_excel(p, sheet_name=None)
    except Exception as e:
        return error_result(KIND, str(e))
    frames = {name: df.dropna(how="all").dropna(axis=1, how="all") for name, df in frames.items()}
    frames = {name: df for name, df in frames.items() if not df.empty}
    if not frames:
        return error_result(KIND, "workbook has no data")
    db_path = build_database(frames, p.stem)
    info = describe_database(db_path)
    props = _sheet_properties(p)
    result = empty_result(KIND)
    result["sql"] = info
    result["text"] = schema_text(info)
    result["meta"] = {"sheets": list(frames.keys()), "sheet_properties": props, "database": str(db_path)}
    if props:
        result["text"] += "\n\nSHEET PROPERTIES:\n" + "\n".join(f"  {k}: {v}" for k, v in props.items())
    total = sum(t["rows"] for t in info["tables"])
    result["summary"] = f"Excel '{p.name}': {len(frames)} sheets, {total} rows total. Loaded into SQLite - use query_sql with db '{db_path.name}'"
    return result