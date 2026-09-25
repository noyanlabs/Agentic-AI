# Reads a .pptx and returns slide text, tables, math (LaTeX), speaker notes and extracted images for the LLM.
from pathlib import Path
import hashlib

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from Interpreter import empty_result, error_result
from Interpreter.math_utils import find_math

KIND = "pptx"
IMAGE_CACHE = Path(__file__).resolve().parent.parent / "LocalStorage" / ".interpreted" / "images"
MAX_IMAGES = 60


def _table_to_rows(table) -> list:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def _walk_shapes(shapes):
    # flattens grouped shapes so nothing nested is missed
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)
        else:
            yield shape


def _save_image(shape, tag: str, stem: str, store: list):
    if len(store) >= MAX_IMAGES:
        return
    try:
        blob, ext = shape.image.blob, shape.image.ext
        digest = hashlib.md5(blob).hexdigest()[:8]
        IMAGE_CACHE.mkdir(parents=True, exist_ok=True)
        out = IMAGE_CACHE / f"{stem}_{tag}_{digest}.{ext}"
        out.write_bytes(blob)
        store.append({"name": out.name, "path": str(out), "mime": f"image/{ext}", "location": tag})
    except Exception:
        pass


def _read_slide(slide, idx: int, stem: str, result: dict):
    parts = []
    for shape in _walk_shapes(slide.shapes):
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text.strip())
        if getattr(shape, "has_table", False) and shape.has_table:
            result["tables"].append({"location": f"slide {idx}", "rows": _table_to_rows(shape.table)})
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            _save_image(shape, f"slide{idx}", stem, result["images"])
        for expr in find_math(shape._element):
            result["math"].append({"location": f"slide {idx}", "latex": expr})
    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        notes = slide.notes_slide.notes_text_frame.text.strip()
    block = f"--- Slide {idx} ---\n" + "\n".join(parts)
    if notes:
        block += f"\n[Speaker notes]: {notes}"
    return block


def initialize(path, slide_range: tuple = None, **options) -> dict:
    p = Path(path)
    if p.suffix.lower() == ".ppt":
        return error_result(KIND, "legacy .ppt is not supported; please convert to .pptx")
    try:
        prs = Presentation(str(p))
    except Exception as e:
        return error_result(KIND, str(e))
    result = empty_result(KIND)
    blocks = []
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, start=1):
        if slide_range and not (slide_range[0] <= i <= slide_range[1]):
            continue
        blocks.append(_read_slide(slide, i, p.stem, result))
    result["text"] = "\n\n".join(blocks)
    result["meta"] = {"slides": total, "tables": len(result["tables"]), "math": len(result["math"]), "images": len(result["images"])}
    result["summary"] = f"PowerPoint '{p.name}': {total} slides, {len(result['tables'])} tables, {len(result['math'])} equations, {len(result['images'])} images"
    return result
